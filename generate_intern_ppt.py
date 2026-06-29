import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Initialize Presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Definitions
COLOR_BLUE_LINE = RGBColor(27, 98, 143)  # Medium steel blue from template
COLOR_BLACK = RGBColor(0, 0, 0)
COLOR_WHITE = RGBColor(255, 255, 255)

# Logo Paths
PES_LOGO_PATH = "pes_logo.png"
PILABS_LOGO_PATH = "pilabs_logo.png"

def add_logos(slide):
    if os.path.exists(PES_LOGO_PATH):
        slide.shapes.add_picture(PES_LOGO_PATH, Inches(0.4), Inches(0.2), width=Inches(1.2), height=Inches(0.8))
    if os.path.exists(PILABS_LOGO_PATH):
        slide.shapes.add_picture(PILABS_LOGO_PATH, Inches(11.73), Inches(0.2), width=Inches(1.2), height=Inches(0.8))

def set_slide_bg_white(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_WHITE

# ----------------- SLIDE 1: TITLE SLIDE -----------------
slide_layout = prs.slide_layouts[6]
slide1 = prs.slides.add_slide(slide_layout)
set_slide_bg_white(slide1)
add_logos(slide1)

header_box = slide1.shapes.add_textbox(Inches(1.5), Inches(0.3), Inches(10.3), Inches(2.2))
tf_header = header_box.text_frame
tf_header.word_wrap = True
tf_header.margin_left = tf_header.margin_top = tf_header.margin_right = tf_header.margin_bottom = 0

p_dept = tf_header.paragraphs[0]
p_dept.text = "DEPARTMENT OF CSE (AI & ML)"
p_dept.font.name = "Arial"
p_dept.font.size = Pt(22)
p_dept.font.bold = True
p_dept.font.color.rgb = COLOR_BLACK
p_dept.alignment = PP_ALIGN.CENTER
p_dept.space_after = Pt(2)

p_sess = tf_header.add_paragraph()
p_sess.text = "Session: June – July 2026"
p_sess.font.name = "Arial"
p_sess.font.size = Pt(12)
p_sess.font.bold = True
p_sess.font.color.rgb = COLOR_BLACK
p_sess.alignment = PP_ALIGN.CENTER
p_sess.space_after = Pt(8)

p_intern = tf_header.add_paragraph()
p_intern.text = "PI Labs Summer Internship"
p_intern.font.name = "Arial"
p_intern.font.size = Pt(22)
p_intern.font.bold = True
p_intern.font.color.rgb = COLOR_BLACK
p_intern.alignment = PP_ALIGN.CENTER
p_intern.space_after = Pt(2)

p_pres = tf_header.add_paragraph()
p_pres.text = "Interim Presentation"
p_pres.font.name = "Arial"
p_pres.font.size = Pt(22)
p_pres.font.bold = True
p_pres.font.color.rgb = COLOR_BLACK
p_pres.alignment = PP_ALIGN.CENTER

title_box = slide1.shapes.add_textbox(Inches(1.5), Inches(2.6), Inches(10.3), Inches(1.4))
tf_title = title_box.text_frame
tf_title.word_wrap = True
tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
p_title = tf_title.paragraphs[0]
p_title.text = "From Conflict to Consensus: Boosting Medical Reasoning via Multi-Round Agentic RAG"
p_title.font.name = "Times New Roman"
p_title.font.size = Pt(32)
p_title.font.bold = True
p_title.font.color.rgb = COLOR_BLACK
p_title.alignment = PP_ALIGN.CENTER

details_box = slide1.shapes.add_textbox(Inches(1.5), Inches(4.3), Inches(10.3), Inches(2.6))
tf_details = details_box.text_frame
tf_details.word_wrap = True
tf_details.margin_left = tf_details.margin_top = tf_details.margin_right = tf_details.margin_bottom = 0

details_text = [
    "Intern1 : Suhas Sreenath   SRN : [SRN]",
    "Intern2 : [Name]           SRN : [SRN]",
    "Intern3 : [Name]           SRN : [SRN]",
    "Intern4 : [Name]           SRN : [SRN]",
    "Mentor: [Mentor Name]"
]
for idx, text in enumerate(details_text):
    p = tf_details.paragraphs[0] if idx == 0 else tf_details.add_paragraph()
    p.text = text
    p.font.name = "Arial"
    p.font.size = Pt(16)
    p.font.color.rgb = COLOR_BLACK
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(4)


# ----------------- HELPER FOR CONTENT SLIDES -----------------
def create_content_slide(title_text):
    slide = prs.slides.add_slide(slide_layout)
    set_slide_bg_white(slide)
    add_logos(slide)
    
    title_box = slide.shapes.add_textbox(Inches(1.8), Inches(0.3), Inches(9.5), Inches(0.6))
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = "Times New Roman"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLACK
    
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.8), Inches(1.0), Inches(11.1), Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_BLUE_LINE
    line.line.fill.background()
    
    content_box = slide.shapes.add_textbox(Inches(1.8), Inches(1.3), Inches(10.5), Inches(5.5))
    tf_content = content_box.text_frame
    tf_content.word_wrap = True
    tf_content.margin_left = tf_content.margin_top = tf_content.margin_right = tf_content.margin_bottom = 0
    
    return slide, tf_content

def add_bullet_points(tf, bullets):
    for idx, text in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = text
        p.font.name = "Arial"
        p.font.size = Pt(18)
        p.font.color.rgb = COLOR_BLACK
        p.level = 0
        p.space_after = Pt(10)


# ----------------- SLIDE 2: PROBLEM STATEMENT -----------------
_, tf_ps = create_content_slide("Problem Statement")
bullets_ps = [
    "Clearly Define the Problem:",
    "  • Large Language Models generate logical-sounding but factually incorrect medical advice due to hallucinations.",
    "  • Standard architectures lack peer-review or clinical debate mechanisms to verify outputs before reaching patients.",
    "  • Vanilla RAG cannot guarantee safety or resolve contradictions when different retrieved journals disagree.",
    "Real-World Relevance:",
    "  • Incorrect AI clinical recommendations can lead to life-threatening errors or wrong diagnoses.",
    "  • Safety-first guardrails are essential to prevent prescribing incorrect medications or dosages.",
    "  • Automatically evaluates consensus and blocks answers when clinical literature is contradictory.",
    "  • Saves healthcare professionals time spent manually cross-referencing conflicting medical sources.",
    "  • Provides reliable, debate-proven consensus answers for clinical decision support systems.",
    "Scope and Constraints:",
    "  • Multi-round debate framework using specialized medical expert personas and dynamic MedCorp retrieval.",
    "  • Extensible design that can easily expand to other high-stakes fields like law and finance.",
    "  • Constraints include high API execution latency, token costs, and vulnerability to server rate limits."
]
add_bullet_points(tf_ps, bullets_ps)


# ----------------- SLIDE 3: LITERATURE REVIEW / BACKGROUND -----------------
_, tf_lr = create_content_slide("Literature Review / Background")
bullets_lr = [
    "Multi-Agent Debate Frameworks: Prior research shows that setting up LLMs to critique each other in a loop improves logical reasoning, but these models still hallucinate clinical facts when not grounded in external evidence.",
    "Medical RAG Systems (e.g., MedRAG): Existing medical QA systems retrieve clinical literature (from the MedCorp corpus) using static queries, but they lack interactive debate and cannot resolve contradictions when source texts disagree.",
    "Knowledge Gaps Identified: Lack of a unified multi-round framework that dynamically analyzes agent disagreements mid-debate to generate targeted search queries to actively retrieve verifying clinical evidence."
]
add_bullet_points(tf_lr, bullets_lr)


# ----------------- SLIDE 4: METHODOLOGY / APPROACH -----------------
_, tf_ma = create_content_slide("Methodology / Approach")
bullets_ma = [
    "Zero-Shot Medical Routing: Evaluates user queries to initiate isolated medical debate sessions.",
    "Multi-Round Persona Debate: Medical Expert A (Primary Care), Expert B (Specialist), and Expert C (Acute Care) generate, critique, and refine answers in parallel rounds.",
    "Active Disagreement Retrieval: A Coordinator agent identifies points of dispute, generates precise search queries, queries MedCorp, and reranks results via MedCPT cross-encoder.",
    "Safety Guardian Layer: Evaluates final consensus, scores clinical confidence (High/Medium/Low), and enforces a safety block if confidence is low."
]
add_bullet_points(tf_ma, bullets_ma)


# ----------------- SLIDE 5: WORK COMPLETED SO FAR -----------------
_, tf_wc = create_content_slide("Work Completed So Far")
bullets_wc = [
    "Retriever Microservice Setup: Deployed local retrieval services running BM25 keyword matching and MedCPT Cross-Encoder reranking.",
    "Medical Agent Pipeline Development: Programmed the 3-expert debate logic, peer critique templates, and moderator synthesis protocols.",
    "Safety Guardian & Consensus Judge: Deployed validation layers to detect factual alignment and audit clinical confidence.",
    "Interactive Testing Interface: Created command-line shells to run and debug multi-round medical debates in real time."
]
add_bullet_points(tf_wc, bullets_wc)


# ----------------- SLIDE 6: PRELIMINARY RESULTS / OBSERVATIONS -----------------
_, tf_pr = create_content_slide("Preliminary Results / Observations")
bullets_pr = [
    "Collaborative Self-Correction: Witnessed Expert B (Specialist) catch and correct an incorrect drug dosage recommendation suggested by Expert A during Round 2 critique.",
    "Guardian Safety Blocking: The Safety Guardian successfully blocked responses to highly ambiguous or dangerous test cases, returning the exact reason for low confidence.",
    "Rate Limit Adjustments: Configured error handling to manage Groq API rate limits (429 errors) during simultaneous agent invocations.",
    "Evaluation Logs: Verified successful test runs on sample MedMCQA multiple-choice questions."
]
add_bullet_points(tf_pr, bullets_pr)


# ----------------- SLIDE 7: PLAN FOR THE REMAINING WEEKS -----------------
_, tf_pl = create_content_slide("Plan for the Remaining Weeks")
bullets_pl = [
    "Batch Evaluation on Dataset: Run the full MedMCQA dataset on the evaluator script (ma_rag_evaluator.py) using Llama-3.1-8B-Instant and Qwen-3-8B.",
    "Analyze Test-Time Scaling: Track and chart how answer accuracy improves as the number of debate rounds and active retrieval loops increases.",
    "Prompt Engineering & Parameter Tuning: Adjust system temperatures and top_k retrieval parameters to optimize answer consensus.",
    "Final Report & Documentation: Compile experimental results and figures, write the final thesis report, and prepare the final presentation."
]
add_bullet_points(tf_pl, bullets_pl)

# Save the Presentation
output_filename = "MA_RAG_Interim_Presentation.pptx"
prs.save(output_filename)
prs.save("ma_rag_presentation.pptx")
print(f"Presentation saved successfully as {output_filename}")
